import csv
import hashlib
import io
import json
import re
import sqlite3
from datetime import date, datetime, timezone
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .core import ROOT, TYPES, TZ, GOOD, Problem, now, uid, pack, parse_day
from .google import SourceError

NATIVE_METRICS = {
    'Google Search - Mobile': 'search_mobile', 'Google Search - Desktop': 'search_desktop',
    'Google Maps - Mobile': 'maps_mobile', 'Google Maps - Desktop': 'maps_desktop',
    'Calls': 'calls', 'Directions': 'directions', 'Website clicks': 'website_clicks',
}
GMB_COLS = ['location', 'name', *NATIVE_METRICS.values()]
LEGACY = {
    'facebook-ads': 'report.html', 'seo': 'seo-report.html', 'facebook-content': 'facebook-content-report.html',
    'facebook-content-30d': 'facebook-content-30D.html', 'facebook-content-6m': 'facebook-content-6M.html',
    'tiktok': 'tiktok-ads-report.html', 'gmb': 'gmb-report.html',
}
LABELS = {
    'activeUsers':'Người dùng hoạt động', 'sessions':'Phiên truy cập', 'screenPageViews':'Lượt xem trang',
    'engagementRate':'Tỷ lệ tương tác', 'clicks':'Lượt nhấp', 'impressions':'Lượt hiển thị',
    'ctr':'CTR', 'position':'Vị trí', 'top5':'Top 5', 'top10':'Top 10', 'top20':'Top 20', 'top100':'Top 100',
    'views':'Lượt hiển thị', 'calls':'Cuộc gọi', 'directions':'Lượt chỉ đường', 'website_clicks':'Nhấp website',
    'date':'Ngày', 'query':'Truy vấn', 'keyword':'Từ khóa', 'url':'URL đích', 'location':'Mã cơ sở', 'name':'Tên cơ sở',
    'search_mobile':'Tìm kiếm trên điện thoại', 'search_desktop':'Tìm kiếm trên máy tính',
    'maps_mobile':'Maps trên điện thoại', 'maps_desktop':'Maps trên máy tính',
}


def format_metric(value, key):
    if value is None:
        return '—'
    if isinstance(value, (int, float)):
        if key in ('ctr', 'engagementRate'):
            return f'{value * 100:.2f}%'.replace('.', ',')
        return f'{value:,.1f}'.rstrip('0').rstrip('.').replace(',', '_').replace('.', ',').replace('_', '.')
    return value


def parse_csv(content, filename, start, end):
    start, end = parse_day(start).isoformat(), parse_day(end).isoformat()
    if start > end:
        raise Problem('Kỳ CSV không hợp lệ.')
    dates = re.findall(r'\d{4}-\d{1,2}-\d{1,2}', filename)
    if len(dates) >= 2:
        try:
            named = [datetime.strptime(d, '%Y-%m-%d').date().isoformat() for d in dates[:2]]
        except ValueError:
            raise Problem('Ngày trong tên file CSV không hợp lệ.') from None
        if named != [start, end]:
            raise Problem('Kỳ đã chọn không khớp kỳ trong tên file CSV. Hãy kiểm tra lại.')
    try:
        decoded = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise Problem('CSV phải dùng mã hóa UTF-8.') from None
    try:
        dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(decoded), dialect=dialect)
    headers = reader.fieldnames or []
    if len(headers) != len(set(headers)):
        raise Problem('File có cột trùng tên.')
    native = 'Store code' in headers
    required = ['Store code', 'Business name', *NATIVE_METRICS] if native else GMB_COLS
    missing = [h for h in required if h not in headers]
    if missing:
        raise Problem('CSV thiếu cột: ' + ', '.join(missing))
    entries, seen = [], set()
    for line, raw in enumerate(reader, start=2):
        if line > 10002:
            raise Problem('CSV vượt giới hạn 10.000 dòng.')
        location = (raw.get('Store code' if native else 'location') or '').strip()
        if native and line == 2 and not location and not (raw.get('Business name') or '').strip():
            # Google exports one explanatory row immediately after the header.
            if any('Number of ' in str(v) for v in raw.values()):
                continue
        if not any(str(v or '').strip() for v in raw.values()):
            continue
        if None in raw or not location:
            raise Problem(f'Dòng {line}: thiếu mã cơ sở hoặc số cột không đúng.')
        if location in seen:
            raise Problem(f'Dòng {line}: mã cơ sở bị trùng trong file.')
        seen.add(location)
        entry = {'location': location, 'name': (raw.get('Business name' if native else 'name') or location).strip()}
        for original, key in NATIVE_METRICS.items():
            value = str(raw.get(original if native else key) or '').strip()
            if not re.fullmatch(r'\d+', value):
                raise Problem(f'Dòng {line}: {original} phải là số nguyên không âm; ô thiếu không được đổi thành 0.')
            number = int(value)
            if number > 10**12:
                raise Problem(f'Dòng {line}: giá trị quá lớn.')
            entry[key] = number
        entries.append(entry)
    if not entries:
        raise Problem('File không có dòng dữ liệu cơ sở hợp lệ.')
    return {'start': start, 'end': end, 'entries': entries}


def save_upload(store, actor, filename, content, start, end):
    data = parse_csv(content, filename, start, end)
    checksum = hashlib.sha256(content).hexdigest()
    upload_id = uid()
    with store.connect(immediate=True) as db:
        if db.execute('SELECT id FROM uploads WHERE sha256=?', (checksum,)).fetchone():
            raise Problem('File này đã được tải lên trước đó.', 409)
        for r in data['entries']:
            overlaps = db.execute('SELECT upload_id FROM upload_locations WHERE location=? AND start_date<=? AND end_date>=?',
                                  (r['location'], data['end'], data['start'])).fetchone()
            if overlaps:
                raise Problem(f"Cơ sở {r['location']} đã có dữ liệu trong kỳ chồng lấn. Chọn kỳ CSV khác.", 409)
        db.execute('INSERT INTO uploads VALUES (?,?,?,?,?,?,?,?)', (upload_id, Path(filename.replace('\\', '/')).name[:200], checksum, actor, data['start'], data['end'], now(), pack(data)))
        for r in data['entries']:
            db.execute('INSERT INTO upload_locations VALUES (?,?,?,?)', (upload_id, r['location'], data['start'], data['end']))
    store.event(actor, 'upload_csv', {'upload_id': upload_id, 'start': start, 'end': end, 'locations': len(data['entries'])})
    return upload_id


def gmb_data(store, p):
    def read(upload_id, start, end):
        row = store.one('SELECT * FROM uploads WHERE id=?', (upload_id,))
        if not row or row['start_date'] != start or row['end_date'] != end:
            raise SourceError('invalid_data', 'File CSV phải khớp chính xác kỳ yêu cầu. Không chia tỷ lệ hoặc suy đoán dữ liệu theo ngày.')
        data = json.loads(row['data'])
        if bool(data.get('demo')) != bool(p.get('demo')):
            raise SourceError('invalid_data', 'Chọn đúng chế độ Demo hoặc dữ liệu thật tương ứng với file CSV.')
        totals = {m: sum(r[m] for r in data['entries']) for m in NATIVE_METRICS.values()}
        totals['views'] = sum(totals[m] for m in ('search_mobile', 'search_desktop', 'maps_mobile', 'maps_desktop'))
        return data['entries'], totals

    entries, totals = read(p['upload_id'], p['start'], p['end'])
    previous_entries, previous = [], None
    if p['compare']:
        previous_entries, previous = read(p['previous_upload_id'], p['previous_start'], p['previous_end'])
        if {r['location'] for r in entries} != {r['location'] for r in previous_entries}:
            raise SourceError('invalid_data', 'Hai kỳ CSV cần có cùng danh sách cơ sở để so sánh.')
    return {'demo': bool(p.get('demo')), 'entries': entries, 'totals': totals, 'previous': previous, 'previous_entries': previous_entries,
            'latest_available_date': p['end'], 'timezone': TZ, 'asset': 'CSV nhập thủ công',
            'warnings': (['DEMO · Số liệu mô phỏng để trình diễn.'] if p.get('demo') else []) + ['Dữ liệu tổng theo kỳ CSV; không có phân bổ theo từng ngày.']}


def valid_dataset(dataset):
    required = TYPES[dataset['params']['report_type']][1]
    return bool(required) and all(dataset['sources'].get(s, {}).get('status') in GOOD for s in required)


def render_report(dataset, settings):
    env = Environment(loader=FileSystemLoader(ROOT / 'templates'), autoescape=select_autoescape())
    env.filters['metric'] = format_metric
    return env.get_template('report.html').render(dataset=dataset, title=TYPES[dataset['params']['report_type']][0], settings=settings, timezone=TZ, labels=LABELS,
                                                 statuses={'ready':'Sẵn sàng','delayed':'Dữ liệu có độ trễ'})


def save_report(store, actor, dataset):
    if not valid_dataset(dataset):
        raise Problem('Nguồn bắt buộc chưa hợp lệ. Chưa thể lưu báo cáo để xuất bản.', 409)
    html = render_report(dataset, dataset.get('organization', {'name': 'KinderHealth', 'author': ''}))
    return insert_report(store, actor, dataset['params']['report_type'], html, dataset['id'],
                         dataset['params']['start'], dataset['params']['end'], True, 'demo' if dataset['params'].get('demo') else 'generated', now())


def insert_report(store, actor, report_type, html, dataset_id, start, end, valid, origin, created_at):
    checksum, report_id = hashlib.sha256(html.encode()).hexdigest(), uid()
    with store.connect(immediate=True) as db:
        existing = db.execute('SELECT id FROM reports WHERE report_type=? AND content_hash=?', (report_type, checksum)).fetchone()
        if existing:
            return existing['id']
        version = db.execute('SELECT COALESCE(MAX(version),0)+1 AS version FROM reports WHERE report_type=?', (report_type,)).fetchone()['version']
        db.execute('INSERT INTO reports VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                   (report_id, report_type, TYPES[report_type][0], version, dataset_id, actor, created_at, start, end, int(valid), origin, checksum, html))
    return report_id


def import_legacy(store, config, actor):
    folder = Path(config['LEGACY_REPORT_DIR']).resolve()
    if not folder.is_dir():
        raise Problem('Chưa tìm thấy thư mục báo cáo cũ ở cấu hình backend.', 404)
    imported = []
    for report_type, name in LEGACY.items():
        path = (folder / name).resolve()
        if path.parent != folder or not path.is_file() or path.stat().st_size > 15 * 1024 * 1024:
            continue
        html = path.read_text(encoding='utf-8-sig')
        created_at = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()
        imported.append(insert_report(store, actor, report_type, html, None, None, None, False, 'legacy', created_at))
    store.event(actor, 'import_legacy', {'reports': len(imported)})
    return imported


def publish(store, actor, report_id):
    row = store.one('SELECT id,report_type,valid,origin FROM reports WHERE id=?', (report_id,))
    if not row:
        raise Problem('Không tìm thấy báo cáo.', 404)
    if row['origin'] == 'demo':
        raise Problem('Báo cáo demo dùng để xem trước và tải xuống; không thay thế bản xuất bản chính thức.', 409)
    if not row['valid']:
        raise Problem('Bản lưu cũ chưa được kiểm chứng dữ liệu. Hãy tạo bản mới từ nguồn hợp lệ trước khi xuất bản.', 409)
    store.execute('INSERT INTO publications VALUES (?,?,?) ON CONFLICT(report_type) DO UPDATE SET report_id=excluded.report_id,published_at=excluded.published_at', (row['report_type'], row['id'], now()))
    store.event(actor, 'publish', {'report_id': report_id, 'report_type': row['report_type']})


def excel_bytes(dataset):
    if not valid_dataset(dataset):
        raise Problem('Không thể xuất Excel: có nguồn bắt buộc lỗi, trống hoặc thiếu kỳ so sánh.', 409)
    wb = Workbook()
    overview = wb.active
    overview.title = 'Tong_quan'
    p = dataset['params']
    metadata = [('Báo cáo', TYPES[p['report_type']][0]), ('Đơn vị', dataset.get('organization', {}).get('name','KinderHealth')), ('Dataset ID', dataset['id']),
                ('Từ ngày', date.fromisoformat(p['start'])), ('Đến ngày', date.fromisoformat(p['end'])),
                ('Múi giờ báo cáo', TZ), ('Lấy dữ liệu (UTC)', dataset['created_at']), ('Xuất file (UTC)', now()),
                ('Loại tìm kiếm GSC', p['search_type']), ('Loại trừ /san-pham/', p['exclude_products'])]
    if p['compare']:
        metadata += [('Kỳ trước từ', date.fromisoformat(p['previous_start'])), ('Kỳ trước đến', date.fromisoformat(p['previous_end']))]
    if p.get('demo'):
        metadata.insert(0, ('DEMO', 'Số liệu mô phỏng để trình diễn, không phải dữ liệu kinh doanh thực tế.'))
    for row in metadata:
        overview.append(row)
    overview.append([])
    overview.append(['Nguồn', 'Trạng thái', 'Dữ liệu đến', 'Lấy dữ liệu UTC', 'Múi giờ nguồn', 'Cảnh báo', 'Tài sản'])
    for source in dataset['sources'].values():
        overview.append([source['label'], source['status'], date.fromisoformat(source['latest_available_date']) if source['latest_available_date'] else None,
                         source['fetched_at'], source.get('timezone'), '\n'.join(source.get('warnings', [])), source.get('asset')])
        for metric, value in source.get('totals', {}).items():
            overview.append([source['label'], metric, value, (source.get('previous') or {}).get(metric)])

    def sheet(name, rows, columns):
        ws = wb.create_sheet(name)
        ws.append([label for key, label in columns])
        for row in rows:
            ws.append([date.fromisoformat(row[key]) if key == 'date' and row.get(key) else row.get(key) for key, label in columns])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    ga = [('date', 'Ngày'), ('activeUsers', 'Người dùng hoạt động'), ('sessions', 'Phiên'), ('screenPageViews', 'Lượt xem trang'), ('engagementRate', 'Tỷ lệ tương tác (0–1)')]
    gsc = [('date', 'Ngày'), ('clicks', 'Lượt nhấp'), ('impressions', 'Lượt hiển thị'), ('ctr', 'CTR (0–1)'), ('position', 'Vị trí trung bình')]
    keyword = [('keyword', 'Từ khóa'), ('position', 'Vị trí'), ('url', 'URL đích'), ('date', 'Ngày tracking')]
    gmb = [(key, key) for key in GMB_COLS]
    for name, source in dataset['sources'].items():
        if name == 'ga4':
            sheet('GA4_theo_ngay', source['daily'], ga)
            sheet('GA4_nguon_truy_cap', source['channels'], [('sessionDefaultChannelGroup', 'Kênh'), ('sessions', 'Phiên')])
            sheet('GA4_trang', source['pages'], [('pagePath', 'Trang'), ('screenPageViews', 'Lượt xem')])
        elif name == 'gsc':
            sheet('GSC_theo_ngay', source['daily'], gsc)
            sheet('GSC_truy_van', source['queries'], [('query', 'Truy vấn'), *gsc[1:]])
        elif name == 'keywords':
            sheet('Keyword_tracking', source['entries'], keyword)
        elif name == 'gmb':
            sheet('GMB_co_so', source['entries'], gmb)
        if p['compare']:
            if name in ('ga4', 'gsc'):
                sheet(name.upper() + '_ky_truoc', source['previous_daily'], ga if name == 'ga4' else gsc)
            elif name in ('keywords', 'gmb'):
                sheet('Keyword_ky_truoc' if name == 'keywords' else 'GMB_ky_truoc', source['previous_entries'], keyword if name == 'keywords' else gmb)
    for ws in wb:
        if p.get('demo'):
            ws.oddHeader.center.text = 'DEMO - Du lieu mo phong'
            ws.oddFooter.center.text = 'DEMO - Khong phai so lieu kinh doanh thuc te'
        for row in ws:
            for cell in row:
                # Force strings to remain text, including untrusted strings beginning with '='.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if isinstance(cell.value, date):
                    cell.number_format = 'dd/mm/yyyy'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='00685F')
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 26 if i < 5 else 42
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
