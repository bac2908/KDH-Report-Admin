import io
import json
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from kdh.app import create_app
from kdh.demo import seed_demo


class DemoTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app({'DATA_DIR': self.temp.name, 'TESTING': True, 'APP_URL': 'http://localhost',
                               'DATABASE_URL':'', 'VERCEL':False, 'JOB_MODE':'worker', 'ENCRYPTION_KEY':'',
                               'INITIAL_ADMIN_EMAIL':'', 'INITIAL_ADMIN_PASSWORD':'',
                               'GOOGLE_CLIENT_ID': '', 'GOOGLE_CLIENT_SECRET': '', 'COOKIE_SECURE': False})
        self.store = self.app.extensions['store']
        self.client = self.app.test_client()
        self.headers = {'X-KDH-Request': '1'}
        account = {'name': 'Demo test', 'email': 'demo@example.test', 'password': 'Test-password-123'}
        self.client.post('/api/auth/setup', json=account, headers=self.headers)
        auth = self.client.post('/api/auth/login', json=account, headers=self.headers).json
        self.actor = auth['user']['id']
        self.headers['X-CSRF-Token'] = auth['csrf']

    def tearDown(self):
        self.temp.cleanup()

    def post(self, url, data):
        return self.client.post(url, json=data, headers=self.headers)

    def test_seed_is_idempotent_preserves_auth_and_labels_artifacts(self):
        users = self.store.all('SELECT * FROM users')
        sessions = self.store.all('SELECT * FROM sessions')
        result = seed_demo(self.store, self.actor)
        self.assertEqual(result, seed_demo(self.store, self.actor))
        self.assertEqual(users, self.store.all('SELECT * FROM users'))
        self.assertEqual(sessions, self.store.all('SELECT * FROM sessions'))
        self.assertEqual(self.store.one('SELECT COUNT(*) AS n FROM reports')['n'], 6)
        self.assertEqual(self.store.one('SELECT COUNT(*) AS n FROM jobs')['n'], 9)
        o = self.client.get('/api/overview').json
        self.assertFalse(o['google_connected'])
        self.assertEqual(o['sources'], {})
        self.assertGreater(o['demo_sources']['ga4']['totals']['sessions'], 1000)
        self.assertEqual(len(o['demo_datasets']), 5)
        for report in self.client.get('/api/reports').json:
            self.assertEqual(report['origin'], 'demo')
            html = self.client.get('/api/reports/'+report['id']+'/html').data.decode()
            self.assertIn('SỐ LIỆU MÔ PHỎNG', html)
            self.assertEqual(self.post('/api/reports/'+report['id']+'/publish', {}).status_code, 409)
        self.assertFalse(self.store.all('SELECT * FROM publications'))
        export = self.client.get('/api/jobs/'+result['export_job_id']+'/download')
        self.assertEqual(export.status_code, 200)
        self.assertIn('DEMO_', export.headers['Content-Disposition'])
        wb = load_workbook(io.BytesIO(export.data))
        self.assertEqual(wb['Tong_quan']['A1'].value, 'DEMO')
        self.assertEqual(wb['GA4_theo_ngay'].max_row, 29)
        self.assertEqual(wb['Keyword_tracking'].max_row, 49)
        self.assertIn('GA4_ky_truoc', wb.sheetnames)
        wb.close()
        export.close()

    def test_explicit_demo_rerun_and_retry_do_not_call_google_real_failure_stays_real(self):
        seed = seed_demo(self.store, self.actor)
        d = self.client.get('/api/datasets/'+seed['dataset_id']).json
        worker = self.app.extensions['worker']
        with patch.object(self.app.extensions['google'], 'fetch', side_effect=AssertionError('No Google in demo')) as fetch:
            job = self.post('/api/analyses', d['params']).json['job_id']
            worker.run_one()
            job = self.client.get('/api/jobs/'+job).json
            self.assertEqual(job['status'], 'succeeded')
            data = self.client.get('/api/datasets/'+job['dataset_id']).json
            self.assertTrue(data['params']['demo'])
            self.assertEqual(data['sources']['ga4']['totals'], d['sources']['ga4']['totals'])
            failed = next(j for j in self.client.get('/api/jobs').json if j['status']=='failed')
            retry = self.post('/api/jobs/'+failed['id']+'/retry', {}).json['job_id']
            worker.run_one()
            self.assertEqual(self.client.get('/api/jobs/'+retry).json['status'], 'succeeded')
            fetch.assert_not_called()
        real = self.post('/api/analyses', {**d['params'], 'demo': False}).json['job_id']
        worker.run_one()
        job = self.client.get('/api/jobs/'+real).json
        self.assertEqual(job['status'], 'failed')
        real_data = self.client.get('/api/datasets/'+job['dataset_id']).json
        self.assertFalse(real_data['params'].get('demo'))
        self.assertFalse(real_data['exportable'])
        self.assertTrue(all('totals' not in s for s in real_data['sources'].values()))

    def test_demo_flag_requires_installation_and_csv_cannot_be_used_as_real_data(self):
        p = {'report_type': 'seo', 'start': '2026-01-01', 'end': '2026-01-28', 'demo': True}
        self.assertEqual(self.post('/api/analyses', p).status_code, 409)
        self.assertEqual(self.post('/api/analyses', {**p, 'demo':'false'}).status_code, 400)
        seed_demo(self.store, self.actor)
        row = self.store.one("SELECT data FROM datasets WHERE report_type='gmb'")
        p = json.loads(row['data'])['params']
        for demo, expected in [(False, 'failed'), (True, 'succeeded')]:
            r = self.post('/api/analyses', {**p, 'demo': demo})
            self.app.extensions['worker'].run_one()
            j = self.client.get('/api/jobs/'+r.json['job_id']).json
            self.assertEqual(j['status'], expected)


if __name__ == '__main__':
    unittest.main()
