"""Exercise the real server/worker at the isolated tests/serve_ui.py port."""
import io
import secrets
import time

import requests
from openpyxl import load_workbook


def main():
    base = 'http://127.0.0.1:8091'
    session = requests.Session()
    session.headers.update({'X-KDH-Request':'1', 'Origin':base})
    email='http-'+secrets.token_hex(4)+'@example.test'
    password=secrets.token_urlsafe(24)

    def request(method,path,**kwargs):
        r=session.request(method,base+path,timeout=30,**kwargs)
        assert r.ok, (r.status_code,r.text[:200])
        return r

    me=request('GET','/api/auth/me').json()
    assert me['setup_required'], 'Use a fresh isolated UI test database.'
    request('POST','/api/auth/setup',json={'name':'HTTP test only','email':email,'password':password})
    response=request('POST','/api/auth/login',json={'email':email,'password':password})
    assert 'HttpOnly' in response.headers['Set-Cookie'] and 'SameSite=Lax' in response.headers['Set-Cookie']
    session.headers['X-CSRF-Token']=response.json()['csrf']
    for path in ['/','/static/app.js','/static/views.js','/static/api.js','/static/ui.js','/static/style.css']:
        r=request('GET',path)
        if path.endswith('.js'):
            assert 'javascript' in r.headers['Content-Type']
    csv=('location,name,search_mobile,search_desktop,maps_mobile,maps_desktop,calls,directions,website_clicks\n'
         'test-only,HTTP test fixture,10,20,30,40,5,8,3\n')
    upload=request('POST','/api/uploads',files={'file':('test-2026-01-01-2026-01-28.csv',csv,'text/csv')},
                   data={'start':'2026-01-01','end':'2026-01-28'}).json()
    params={'report_type':'gmb','start':'2026-01-01','end':'2026-01-28','upload_id':upload['upload_id']}

    def wait(job_id):
        for _ in range(30):
            job=request('GET','/api/jobs/'+job_id).json()
            if job['status'] not in ('queued','running'):
                assert job['status']=='succeeded', job
                return job
            time.sleep(.3)
        raise AssertionError('Worker timeout')

    job=wait(request('POST','/api/analyses',json=params).json()['job_id'])
    dataset=request('GET','/api/datasets/'+job['dataset_id']).json()
    assert dataset['sources']['gmb']['totals']['views']==100
    export=wait(request('POST','/api/datasets/'+dataset['id']+'/export',json={}).json()['job_id'])
    xlsx=request('GET','/api/jobs/'+export['id']+'/download')
    wb=load_workbook(io.BytesIO(xlsx.content))
    assert 'GMB_co_so' in wb.sheetnames
    assert wb['GMB_co_so']['C2'].value==10
    wb.close()
    rid=request('POST','/api/datasets/'+dataset['id']+'/save',json={}).json()['report_id']
    preview=request('GET','/api/reports/'+rid+'/html')
    assert 'sandbox allow-scripts' in preview.headers['Content-Security-Policy']
    request('POST','/api/reports/'+rid+'/publish',json={})
    reports=request('GET','/api/reports').json()
    assert any(r['id']==rid and r['published_at'] for r in reports)
    request('POST','/api/auth/logout',json={})
    assert session.get(base+'/api/reports',timeout=5).status_code==401
    session.close()
    print('HTTP smoke passed: login, static modules, CSV, worker, Excel, HTML, publication, logout.')


if __name__=='__main__':
    main()
