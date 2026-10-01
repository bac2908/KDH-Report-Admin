import io
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from kdh.app import create_app
from kdh.core import filters, pack, now, TYPES
from kdh.google import SourceError, parse_keywords
from kdh.jobs import next_run, period_dates


def fixture(source, p):
    common = {'latest_available_date': p['end'], 'timezone': 'Asia/Ho_Chi_Minh', 'asset': source, 'warnings': [], 'previous': None}
    if source == 'ga4':
        totals = {'activeUsers': 21, 'sessions': 35, 'screenPageViews': 49, 'engagementRate': .5}
        return {**common, 'totals': totals, 'daily': [{'date': p['end'], **totals}], 'channels': [], 'pages': [], 'previous_daily': []}
    if source == 'gsc':
        totals = {'clicks': 4, 'impressions': 50, 'ctr': .08, 'position': 3.5}
        return {**common, 'totals': totals, 'daily': [{'date': p['end'], **totals}], 'queries': [{'query': '=HYPERLINK("https://evil.invalid")', **totals}], 'previous_daily': []}
    return {**common, 'totals': {'top5': 1, 'top10': 1, 'top20': 1, 'top100': 1}, 'entries': [{'keyword': 'phòng khám nhi', 'date':p['end'], 'position':4, 'url':'https://kinderhealth.vn/'}], 'previous_entries': []}


class AppTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        # Tests own their OAuth configuration even after real credentials are added to .env.
        self.app = create_app({'DATA_DIR':self.temp.name, 'TESTING':True, 'APP_URL':'http://localhost',
                               'GOOGLE_CLIENT_ID':'', 'GOOGLE_CLIENT_SECRET':'', 'COOKIE_SECURE':False,
                               'GOOGLE_REDIRECT_URI':'http://localhost/api/google/callback',
                               'LEGACY_REPORT_DIR':str(Path(self.temp.name)/'legacy')})
        self.store = self.app.extensions['store']
        self.google = self.app.extensions['google']
        self.worker = self.app.extensions['worker']
        self.client = self.app.test_client()
        self.headers = {'X-KDH-Request':'1'}
        result = self.client.post('/api/auth/setup', json={'name':'Test Admin','email':'admin@example.test','password':'Test-password-123'},headers=self.headers)
        self.assertEqual(result.status_code,201,result.json)
        result = self.client.post('/api/auth/login',json={'email':'admin@example.test','password':'Test-password-123'},headers=self.headers)
        self.assertEqual(result.status_code,200,result.json)
        self.headers['X-CSRF-Token']=result.json['csrf']
        self.admin=result.json['user']

    def tearDown(self):
        self.temp.cleanup()

    def post(self,url, data=None):
        return self.client.post(url,json=data or {},headers=self.headers)

    def analyze(self, kind='seo', extra=None, fetch=fixture):
        p={'report_type':kind,'start':'2026-01-01','end':'2026-01-28','compare':False,**(extra or {})}
        r=self.post('/api/analyses',p)
        self.assertEqual(r.status_code,202,r.json)
        with patch.object(self.google,'fetch',side_effect=fetch):
            self.worker.run_one()
        job=self.client.get('/api/jobs/'+r.json['job_id']).json
        return job, self.client.get('/api/datasets/'+job['dataset_id']).json if job['dataset_id'] else None

    def test_session_csrf_and_local_bootstrap(self):
        self.assertEqual(self.post('/api/auth/setup').status_code,403)
        self.assertEqual(self.client.post('/api/settings',json={},headers={'X-KDH-Request':'1'}).status_code,403)
        self.assertEqual(self.client.patch('/api/settings',json={'name':'x'},headers={**self.headers,'Origin':'https://evil.invalid'}).status_code,403)
        self.post('/api/auth/logout')
        self.assertEqual(self.client.get('/api/reports').status_code,401)

    def test_viewer_acl_and_session_revocation(self):
        job,data=self.analyze()
        r=self.post('/api/users',{'name':'Viewer','email':'viewer@example.test','password':'Viewer-password-123','role':'viewer','allowed':['gsc']})
        viewer_id=r.json['id']
        viewer=self.app.test_client()
        session=viewer.post('/api/auth/login',json={'email':'viewer@example.test','password':'Viewer-password-123'},headers={'X-KDH-Request':'1'}).json
        h={'X-KDH-Request':'1','X-CSRF-Token':session['csrf']}
        self.assertEqual(viewer.get('/api/google').status_code,403)
        self.assertEqual(viewer.get('/api/users').status_code,403)
        self.assertEqual(viewer.get('/api/datasets/'+data['id']).status_code,403)
        self.assertEqual(viewer.get('/api/jobs/'+job['id']).status_code,403)
        self.assertEqual(viewer.post('/api/datasets/'+data['id']+'/save',json={},headers=h).status_code,403)
        self.assertEqual(viewer.post('/api/analyses',json={'report_type':'ga4','start':'2026-01-01','end':'2026-01-28'},headers=h).status_code,403)
        r=self.client.patch('/api/users/'+viewer_id,json={'active':False},headers=self.headers)
        self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(viewer.get('/api/overview').status_code,401)

    def test_duplicate_jobs_and_date_validation(self):
        p={'report_type':'ga4','start':'2026-01-01','end':'2026-01-28'}
        first=self.post('/api/analyses',p); second=self.post('/api/analyses',p)
        self.assertEqual(first.json['job_id'],second.json['job_id'])
        self.assertFalse(second.json['created'])
        self.assertEqual(self.post('/api/analyses',{**p,'start':'2026-02-01'}).status_code,400)
        self.assertEqual(self.post('/api/analyses',{**p,'end':'2099-01-01'}).status_code,400)
        self.assertEqual(self.post('/api/analyses',{**p,'start':'01/01/2026'}).status_code,400)

    def test_partial_failure_no_export_no_fake_data_and_retry(self):
        def fetch(source,p):
            if source=='gsc':raise SourceError('permission_denied','Thiếu quyền')
            return fixture(source,p)
        job,data=self.analyze(fetch=fetch)
        self.assertEqual(job['status'],'partial')
        self.assertEqual(data['sources']['gsc']['status'],'permission_denied')
        self.assertNotIn('totals',data['sources']['gsc'])
        self.assertFalse(data['exportable'])
        self.assertEqual(self.post('/api/datasets/'+data['id']+'/export').status_code,409)
        self.assertEqual(self.post('/api/datasets/'+data['id']+'/save').status_code,409)
        r=self.post('/api/jobs/'+job['id']+'/retry')
        self.assertEqual(r.status_code,202)
        with patch.object(self.google,'fetch',side_effect=fixture):self.worker.run_one()
        retried=self.client.get('/api/jobs/'+r.json['job_id']).json
        self.assertEqual(retried['status'],'succeeded')
        self.assertNotEqual(retried['dataset_id'],data['id'])

    def test_export_uses_snapshot_dates_numbers_and_no_formulas(self):
        job,data=self.analyze()
        self.assertEqual(job['status'],'succeeded')
        result=self.post('/api/datasets/'+data['id']+'/export')
        export_id=result.json['job_id']
        with patch.object(self.google,'fetch',side_effect=AssertionError('Export must never fetch data')):
            self.worker.run_one()
        download=self.client.get('/api/jobs/'+export_id+'/download')
        self.assertEqual(download.status_code,200)
        self.assertIn('KinderHealth_Report_2026-01-01_2026-01-28.xlsx',download.headers['Content-Disposition'])
        wb=load_workbook(io.BytesIO(download.data))
        for sheet in ['Tong_quan','GA4_theo_ngay','GSC_theo_ngay','GSC_truy_van','Keyword_tracking']:
            self.assertIn(sheet,wb.sheetnames)
        self.assertIsInstance(wb['GA4_theo_ngay']['A2'].value,datetime)
        self.assertEqual(wb['GA4_theo_ngay']['B2'].value,21)
        self.assertEqual(wb['GSC_truy_van']['A2'].data_type,'s')
        self.assertTrue(wb['GSC_truy_van']['A2'].value.startswith('='))
        self.assertIn(data['id'],str(list(wb['Tong_quan'].values)))
        download.close()
        wb.close()

    def test_lag_preserves_requested_period_and_last_good_date(self):
        def delayed(source,p):
            r=fixture(source,p);r['latest_available_date']='2026-01-26';r['daily'][0]['date']='2026-01-26';return r
        job,data=self.analyze('ga4',fetch=delayed)
        self.assertEqual(data['params']['end'],'2026-01-28')
        self.assertEqual(data['sources']['ga4']['status'],'delayed')
        self.assertTrue(data['exportable'])
        self.assertEqual(len(data['sources']['ga4']['daily']),1)
        def fail(source,p):raise SourceError('timeout','Quá thời gian')
        job2,data2=self.analyze('ga4',fetch=fail)
        overview=self.client.get('/api/overview').json
        self.assertEqual(overview['sources']['ga4']['last_success_data_date'],'2026-01-26')

    def test_comparison_failure_blocks_export(self):
        job,data=self.analyze('ga4',extra={'compare':True})
        self.assertEqual(data['params']['previous_start'],'2025-12-04')
        self.assertEqual(data['params']['previous_end'],'2025-12-31')
        self.assertEqual(data['sources']['ga4']['status'],'incomplete')
        self.assertFalse(data['exportable'])

    def test_versions_publication_restore_and_failure_preservation(self):
        job,data=self.analyze('ga4')
        first=self.post('/api/datasets/'+data['id']+'/save')
        self.assertEqual(first.status_code,201,first.json)
        rid=first.json['report_id']
        self.assertEqual(self.post('/api/datasets/'+data['id']+'/save').json['report_id'],rid)
        self.assertEqual(self.post('/api/reports/'+rid+'/publish').status_code,200)
        job2,data2=self.analyze('ga4',extra={'end':'2026-01-27'})
        rid2=self.post('/api/datasets/'+data2['id']+'/save').json['report_id']
        self.post('/api/reports/'+rid2+'/publish')
        self.post('/api/reports/'+rid+'/publish')
        self.assertEqual(self.store.one('SELECT report_id FROM publications')['report_id'],rid)
        response=self.client.get('/api/reports/'+rid+'/html')
        self.assertIn('sandbox allow-scripts',response.headers['Content-Security-Policy'])
        self.assertNotIn('allow-same-origin',response.headers['Content-Security-Policy'])

    def test_legacy_import_is_idempotent_unverified_and_read_only(self):
        folder=Path(self.app.config['LEGACY_REPORT_DIR']);folder.mkdir()
        html='<html><h1>Old report</h1><script>window.parent.document</script></html>'
        path=folder/'seo-report.html';path.write_text(html)
        first=self.post('/api/reports/import').json['report_ids']
        second=self.post('/api/reports/import').json['report_ids']
        self.assertEqual(first,second)
        self.assertEqual(path.read_text(),html)
        self.assertEqual(self.post('/api/reports/'+first[0]+'/publish').status_code,409)
        row=self.client.get('/api/reports').json[0]
        self.assertIsNone(row['start_date'])
        self.assertFalse(row['valid'])

    def test_csv_validation_duplicate_overlap_and_gmb_report(self):
        header='location,name,search_mobile,search_desktop,maps_mobile,maps_desktop,calls,directions,website_clicks\n'
        def upload(data,name='gmb.csv',start='2026-01-01',end='2026-01-28'):
            return self.client.post('/api/uploads',data={'file':(io.BytesIO(data.encode()),name),'start':start,'end':end},headers=self.headers)
        data=header+'branch1,Clinic,10,20,30,40,5,8,3\n'
        r=upload(data);self.assertEqual(r.status_code,201,r.json)
        self.assertEqual(upload(data).status_code,409)
        self.assertEqual(upload(data.replace(',10,',',11,')).status_code,409)
        self.assertEqual(upload(header+'branch2,Clinic,,20,30,40,5,8,3\n').status_code,400)
        self.assertEqual(upload(data,'gmb 2026-02-01 2026-02-28.csv').status_code,400)
        job,result=self.analyze('gmb',{'upload_id':r.json['upload_id']})
        self.assertEqual(job['status'],'succeeded')
        self.assertEqual(result['sources']['gmb']['totals']['views'],100)
        self.assertNotIn('daily',result['sources']['gmb'])
        job,result=self.analyze('gmb',{'upload_id':r.json['upload_id'],'end':'2026-01-27'})
        self.assertEqual(result['sources']['gmb']['status'],'invalid_data')

    def test_missing_google_no_demo_and_secrets_not_in_response(self):
        self.assertEqual(self.post('/api/google/connect').status_code,503)
        r=self.post('/api/analyses',{'report_type':'seo','start':'2026-01-01','end':'2026-01-28'})
        self.worker.run_one()
        job=self.client.get('/api/jobs/'+r.json['job_id']).json
        self.assertEqual(job['status'],'failed')
        data=self.client.get('/api/datasets/'+job['dataset_id']).json
        self.assertTrue(all(s['status']=='disconnected' for s in data['sources'].values()))
        self.google.save({'refresh_token':'DO_NOT_EXPOSE','access_token':'DO_NOT_EXPOSE','email':'google@example.test','connected_at':now(),'expires_at':'2099-01-01T00:00:00+00:00'})
        raw=self.client.get('/api/google').data
        self.assertNotIn(b'DO_NOT_EXPOSE',raw)
        self.assertNotIn(b'DO_NOT_EXPOSE',self.store.one('SELECT value FROM secrets')['value'])

    def test_oauth_state_bound_to_session_and_one_time(self):
        self.app.config.update(GOOGLE_CLIENT_ID='client',GOOGLE_CLIENT_SECRET='secret')
        from urllib.parse import urlparse,parse_qs
        url=self.post('/api/google/connect').json['url']
        state=parse_qs(urlparse(url).query)['state'][0]
        record=self.store.one('SELECT * FROM oauth_states')
        with self.assertRaises(Exception):self.google.exchange(state,'code','wrong-session')
        self.assertIsNotNone(self.store.one('SELECT * FROM oauth_states'))
        with patch('kdh.google.call',side_effect=[{'access_token':'access','refresh_token':'refresh','expires_in':3600},{'email':'g@example.test','email_verified':True}]):
            self.google.exchange(state,'code',record['session_id'])
        self.assertIsNone(self.store.one('SELECT * FROM oauth_states'))
        with self.assertRaises(Exception):self.google.exchange(state,'code',record['session_id'])

    def test_scheduler_claims_once_and_disabled_schedule_stays_idle(self):
        r=self.post('/api/schedules',{'name':'Test schedule','report_type':'ga4','frequency':'daily','run_time':'07:00','period':'last7'})
        self.assertEqual(r.status_code,201,r.json)
        schedule_id=r.json['id']
        self.store.execute('UPDATE schedules SET next_run=? WHERE id=?',('2000-01-01T00:00:00+00:00',schedule_id))
        self.worker.tick_schedules();self.worker.tick_schedules()
        self.assertEqual(len(self.store.all('SELECT * FROM jobs')),1)
        with patch.object(self.google,'fetch',side_effect=fixture):self.worker.run_one()
        self.assertEqual(len(self.store.all('SELECT * FROM reports')),1)
        self.assertEqual(len(self.store.all('SELECT * FROM publications')),0)
        self.client.patch('/api/schedules/'+schedule_id,json={'enabled':False},headers=self.headers)
        self.store.execute('UPDATE schedules SET next_run=? WHERE id=?',('2000-01-01T00:00:00+00:00',schedule_id))
        self.worker.tick_schedules()
        self.assertEqual(len(self.store.all('SELECT * FROM jobs')),1)

    def test_keyword_parsing_never_invents_year(self):
        rows=[['Từ khóa','URL','Jul/16'],['test','https://example.test','4']]
        with self.assertRaises(SourceError):parse_keywords(rows)
        entries,_=parse_keywords(rows,'2025')
        self.assertEqual(entries[0]['date'],'2025-07-16')
        entries,_=parse_keywords([['Từ khóa','Vị trí','Ngày'],['test','4','16/07/2026']])
        self.assertEqual(entries[0]['date'],'2026-07-16')
        with self.assertRaises(SourceError):parse_keywords([['Từ khóa','Vị trí','Ngày'],['test','NaN','16/07/2026']])

    def test_schedule_timezone_and_previous_month(self):
        at=datetime(2026,1,31,18,0,tzinfo=timezone.utc) # 01:00 Feb 1 in Vietnam
        self.assertEqual(period_dates('previous_month',at),('2026-01-01','2026-01-31'))
        self.assertEqual(next_run({'run_time':'07:00','frequency':'monthly','weekday':0,'monthday':1},at),'2026-02-01T00:00:00+00:00')

    def test_ga4_queries_distinct_totals_daily_and_orders_top_pages(self):
        p=filters({'report_type':'ga4','start':'2026-01-01','end':'2026-01-28'})
        calls=[]
        def provider(method,url,**kwargs):
            payload=kwargs['json'];calls.append(payload)
            dims=[d['name'] for d in payload.get('dimensions',[])]
            metrics=[m['name'] for m in payload['metrics']]
            row={'dimensionValues':[{'value':'20260128' if d=='date' else 'Direct'} for d in dims],
                 'metricValues':[{'value':'.5' if m=='engagementRate' else '21'} for m in metrics]}
            return {'rows':[row],'metadata':{'timeZone':'Asia/Ho_Chi_Minh'}}
        with patch('kdh.google.call',side_effect=provider):
            result=self.google.ga4(p,{'Authorization':'Bearer test-only'})
        self.assertNotIn('dimensions',calls[0])
        self.assertEqual(result['totals']['activeUsers'],21)
        self.assertEqual(result['daily'][0]['date'],'2026-01-28')
        self.assertEqual(calls[-1]['orderBys'],[{'metric':{'metricName':'screenPageViews'},'desc':True}])

    def test_gsc_filters_applied_to_all_queries_and_metadata(self):
        p=filters({'report_type':'gsc','start':'2026-01-01','end':'2026-01-28','exclude_products':True,'compare':True})
        calls=[]
        def provider(method,url,**kwargs):
            payload=kwargs['json'];calls.append(payload)
            dimensions=payload.get('dimensions',[])
            return {'rows':[{'keys':['2026-01-26' if d=='date' else 'clinic' for d in dimensions],
                             'clicks':3,'impressions':40,'ctr':.075,'position':4}]}
        with patch('kdh.google.call',side_effect=provider):
            result=self.google.gsc(p,{})
        self.assertEqual(len(calls),5)
        self.assertTrue(all(c['dataState']=='final' and c['type']=='web' for c in calls))
        self.assertTrue(all(c['dimensionFilterGroups'][0]['filters'][0]['expression']=='/san-pham/' for c in calls))
        self.assertEqual(result['timezone'],'America/Los_Angeles')
        self.assertEqual(result['latest_available_date'],'2026-01-26')

    def test_native_csv_explanatory_row_and_blank_metrics(self):
        from kdh.reports import parse_csv, NATIVE_METRICS
        header='Store code,Business name,'+','.join(NATIVE_METRICS)+'\n'
        description=',,'+','.join(['Number of interactions']*7)+'\n'
        row='00012,Clinic,10,20,30,40,5,8,3\n'
        result=parse_csv((header+description+row).encode(),'google.csv','2026-01-01','2026-01-28')
        self.assertEqual(result['entries'][0]['location'],'00012')
        with self.assertRaises(Exception):
            parse_csv((header+description+row.replace(',10,',',,')).encode(),'google.csv','2026-01-01','2026-01-28')

    def test_revoked_permissions_before_queue_execution(self):
        r=self.post('/api/analyses',{'report_type':'ga4','start':'2026-01-01','end':'2026-01-28'})
        self.store.execute('UPDATE users SET active=0 WHERE id=?',(self.admin['id'],))
        with patch.object(self.google,'fetch') as fetch:
            self.worker.run_one()
            fetch.assert_not_called()
        row=self.store.one('SELECT * FROM jobs WHERE id=?',(r.json['job_id'],))
        self.assertEqual(row['status'],'failed')
        self.assertIsNone(row['dataset_id'])

    def test_worker_restart_marks_interruption_and_keeps_queued_work(self):
        first=self.post('/api/analyses',{'report_type':'ga4','start':'2026-01-01','end':'2026-01-28'}).json['job_id']
        second=self.post('/api/analyses',{'report_type':'gsc','start':'2026-01-01','end':'2026-01-28'}).json['job_id']
        self.store.execute("UPDATE jobs SET status='running' WHERE id=?",(first,))
        self.worker.stopping.set()
        self.worker.start()
        self.worker.thread.join(timeout=1)
        self.assertEqual(self.store.one('SELECT status FROM jobs WHERE id=?',(first,))['status'],'interrupted')
        self.assertEqual(self.store.one('SELECT status FROM jobs WHERE id=?',(second,))['status'],'queued')


if __name__=='__main__':
    unittest.main()
