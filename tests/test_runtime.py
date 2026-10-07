"""Runtime and optional request-processing behavior using isolated storage."""
import io
import os
import tempfile
import time
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from cryptography.fernet import Fernet
from openpyxl import load_workbook

from kdh.app import create_app
from kdh.core import load_config, now
from kdh.google import REQUEST_DEADLINE, SourceError, call
from test_app import fixture


class RequestModeTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.config = {
            'DATA_DIR': self.temp.name, 'DATABASE_URL': '',
            'JOB_MODE': 'request', 'TESTING': True, 'APP_URL': 'https://kdh.example.test',
            'COOKIE_SECURE': True, 'ENCRYPTION_KEY': Fernet.generate_key().decode(),
            'INITIAL_ADMIN_EMAIL': 'cloud@example.test', 'INITIAL_ADMIN_PASSWORD': 'Cloud-test-password-123',
            'GOOGLE_CLIENT_ID': '', 'GOOGLE_CLIENT_SECRET': '', 'CRON_SECRET': 'test-cron-secret',
        }
        self.app = create_app(self.config)
        self.client = self.app.test_client()
        result = self.client.post('/api/auth/login', json={
            'email': self.config['INITIAL_ADMIN_EMAIL'], 'password': self.config['INITIAL_ADMIN_PASSWORD'],
        }, headers={'X-KDH-Request': '1'})
        self.assertEqual(result.status_code, 200, result.json)
        self.actor = result.json['user']['id']
        self.headers = {'X-KDH-Request': '1', 'X-CSRF-Token': result.json['csrf']}
        self.store = self.app.extensions['store']
        self.worker = self.app.extensions['worker']

    def post(self, path, data=None):
        return self.client.post(path, json=data or {}, headers=self.headers)

    def enqueue(self):
        result = self.post('/api/analyses', {'report_type': 'seo', 'start': '2026-01-01', 'end': '2026-01-28'})
        self.assertEqual(result.status_code, 202, result.json)
        return result.json['job_id']

    def test_no_thread_request_analysis_export_and_idempotent_process(self):
        self.worker.start()
        self.assertIsNone(self.worker.thread)
        job_id = self.enqueue()
        path = '/api/jobs/' + job_id + '/process'
        self.assertEqual(self.client.post(path, json={}).status_code, 403)
        with patch.object(self.app.extensions['google'], 'fetch', side_effect=fixture) as fetch:
            result = self.post(path)
            self.assertEqual(result.json['status'], 'succeeded', result.json)
            self.assertEqual(self.post(path).json['dataset_id'], result.json['dataset_id'])
            self.assertEqual(fetch.call_count, 3)
        dataset_id = result.json['dataset_id']
        export = self.post('/api/datasets/' + dataset_id + '/export').json['job_id']
        self.assertEqual(self.post('/api/jobs/' + export + '/process').json['status'], 'succeeded')
        data = self.client.get('/api/jobs/' + export + '/download')
        self.assertEqual(data.status_code, 200)
        wb = load_workbook(io.BytesIO(data.data))
        self.assertIn('Tong_quan', wb.sheetnames)
        wb.close()
        self.assertFalse((self.store.folder / 'exports').exists())

    def test_cron_auth_executes_due_schedule_without_browser(self):
        self.assertEqual(self.client.get('/api/cron').status_code, 401)
        self.assertEqual(self.client.get('/api/cron', headers={'Authorization': 'Bearer wrong'}).status_code, 401)
        result = self.post('/api/schedules', {'name': 'Cloud daily', 'report_type': 'ga4', 'frequency': 'daily',
                                             'run_time': '07:00', 'period': 'last7', 'compare': False})
        self.assertEqual(result.status_code, 201, result.json)
        self.store.execute('UPDATE schedules SET next_run=?', ('2026-01-01T00:00:00+00:00',))
        with patch.object(self.app.extensions['google'], 'fetch', side_effect=fixture):
            result = self.client.get('/api/cron', headers={'Authorization': 'Bearer test-cron-secret'})
        self.assertEqual(result.status_code, 200, result.json)
        self.assertEqual(result.json, {'processed': 1, 'queued': 0})
        self.assertEqual(len(self.store.all('SELECT * FROM reports')), 1)
        again = self.client.get('/api/cron', headers={'Authorization': 'Bearer test-cron-secret'})
        self.assertEqual(again.json['processed'], 0)

    def test_cold_start_preserves_running_and_recovers_only_stale_jobs(self):
        job_id = self.enqueue()
        self.store.execute("UPDATE jobs SET status='running',started_at=? WHERE id=?", (now(), job_id))
        create_app(self.config)
        self.assertEqual(self.client.get('/api/jobs/' + job_id).json['status'], 'running')
        old = (datetime.now(timezone.utc) - timedelta(minutes=7)).isoformat()
        self.store.execute('UPDATE jobs SET started_at=? WHERE id=?', (old, job_id))
        self.assertEqual(self.client.get('/api/jobs/' + job_id).json['status'], 'interrupted')
        retry = self.post('/api/jobs/' + job_id + '/retry')
        self.assertEqual(retry.status_code, 202)

    def test_demo_seed_is_available_to_admin_and_idempotent(self):
        seed = self.post('/api/demo/seed')
        self.assertEqual(seed.status_code, 200, seed.json)
        self.assertEqual(self.post('/api/demo/seed').json, seed.json)
        self.assertEqual(len(self.store.all('SELECT * FROM reports')), 6)
        self.store.execute("UPDATE users SET role='viewer' WHERE id=?", (self.actor,))
        self.assertEqual(self.post('/api/demo/seed').status_code, 403)

    def test_local_worker_persists_login_and_encrypted_connection_after_restart(self):
        config = {**self.config, 'JOB_MODE': 'worker', 'ENCRYPTION_KEY': '',
                  'APP_URL': 'http://localhost', 'COOKIE_SECURE': False}
        first = create_app(config)
        first.extensions['google'].save({'access_token': 'runtime-test-token'})
        key = (self.store.folder / 'encryption.key').read_bytes()
        restarted = create_app(config)
        self.assertEqual(restarted.extensions['google'].read()['access_token'], 'runtime-test-token')
        self.assertEqual((self.store.folder / 'encryption.key').read_bytes(), key)
        login = restarted.test_client().post('/api/auth/login', json={
            'email': config['INITIAL_ADMIN_EMAIL'], 'password': config['INITIAL_ADMIN_PASSWORD'],
        }, headers={'X-KDH-Request': '1'})
        self.assertEqual(login.status_code, 200, login.json)

    def test_expired_google_budget_never_calls_provider(self):
        token = REQUEST_DEADLINE.set(time.monotonic() + 30)
        try:
            with patch('kdh.google.requests.request') as request:
                with self.assertRaises(SourceError):
                    call('GET', 'https://example.test')
                request.assert_not_called()
        finally:
            REQUEST_DEADLINE.reset(token)

    def test_default_runtime_uses_worker_and_persistent_storage(self):
        with patch.dict(os.environ, {}, clear=True), patch('kdh.core.ROOT', Path(self.temp.name)):
            config = load_config()
        self.assertEqual(config['JOB_MODE'], 'worker')
        self.assertEqual(config['DATA_DIR'], str(Path(self.temp.name) / 'instance'))
        self.assertEqual(config['DATABASE_URL'], '')
        self.assertFalse(config['COOKIE_SECURE'])
        self.assertEqual(config['MAX_CONTENT_LENGTH'], 5 * 1024 * 1024)
        with self.assertRaisesRegex(RuntimeError, 'JOB_MODE'):
            create_app({**self.config, 'JOB_MODE': 'invalid'})

    def test_large_response_preserved_with_security_headers(self):
        content = b'x' * (4 * 1024 * 1024 + 1)
        with self.app.test_request_context('/'):
            response = self.app.process_response(self.app.response_class(content))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, content)
        self.assertEqual(response.headers['X-Content-Type-Options'], 'nosniff')
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
